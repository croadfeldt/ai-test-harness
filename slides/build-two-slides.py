"""Build the two slides as an editable PowerPoint file, the same content as two-slides.md.

    python slides/build-two-slides.py [out.pptx]      default: harness/out/two-slides.pptx

Needs python-pptx. Every box is a shape with text, so the deck owner can restyle it in place.
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

RED, INK, GREY, LINE, MUTED, WHITE, PINK, GREEN, MID = (RGBColor(0xEE, 0, 0), RGBColor(0x15, 0x15, 0x15), RGBColor(0xF0, 0xF0, 0xF0),
                                                      RGBColor(0xC7, 0xC7, 0xC7), RGBColor(0x4F, 0x52, 0x55), RGBColor(0xFF, 0xFF, 0xFF),
                                                      RGBColor(0xFF, 0xD6, 0xD6), RGBColor(0x3E, 0x86, 0x35), RGBColor(0xD2, 0xD2, 0xD2))
FONT = "Red Hat Text"
COMIC = Path(__file__).resolve().parent / "xkcd-1172-workflow.png"
L = 0.6           # left margin, inches
W = 13.333 - 2 * L


def box(slide, x, y, w, h, lines, *, fill=GREY, line=LINE, line_w=0.75, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, inset=0.07):
    """A rectangle with text. `lines` is a list of (text, size, bold, color, italic) tuples, one paragraph each."""
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(line_w)
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(inset); tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, (text, size, bold, color, italic) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = text
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic; r.font.color.rgb = color; r.font.name = FONT
    return s


def text(slide, x, y, w, h, paragraphs, *, size=16, color=INK, align=PP_ALIGN.LEFT, space_after=10, anchor=MSO_ANCHOR.TOP):
    t = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = t.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space_after)
        runs = para if isinstance(para, list) else [(para, False)]
        for s, bold in runs:
            r = p.add_run(); r.text = s; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = FONT
    return t


def title(slide, s):
    text(slide, L, 0.35, W, 0.7, [s], size=32, color=RED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def band(slide, x, y, w, label, *, red):
    box(slide, x, y, w, 0.3, [(label.upper(), 11, True, WHITE if red else INK, False)], fill=RED if red else MID, line=None, align=PP_ALIGN.CENTER)


def grid(slide, x, y, w, rows, *, col_w=None, row_h=0.3, size=10, header=True, upper=True, last_right=False):
    """Rows of cells; each cell is a string or a list of (text, size, bold, color, italic). Returns the height used."""
    n = len(rows[0])
    widths = col_w or [w / n] * n
    cy = y
    for ri, row in enumerate(rows):
        cx = x
        for ci, cell in enumerate(row):
            lines = cell if isinstance(cell, list) else [(cell.upper() if upper else cell, size, header and ri == 0, INK, False)]
            align = PP_ALIGN.RIGHT if (last_right and ci == n - 1) else PP_ALIGN.LEFT
            box(slide, cx, cy, widths[ci], row_h, lines, align=align, anchor=MSO_ANCHOR.TOP)
            cx += widths[ci]
        cy += row_h
    return cy - y


def down(slide, x, y, w):
    text(slide, x, y, w, 0.18, ["▼"], size=10, color=MUTED, align=PP_ALIGN.CENTER, space_after=0)


def slide_one(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "AI Test Harness")
    # the comic, with its source line
    s.shapes.add_picture(str(COMIC), Inches(L), Inches(1.3), width=Inches(2.2))
    text(s, L, 4.42, 2.2, 0.5, ["“Workflow” by Randall Munroe, xkcd.com/1172. Licensed CC BY-NC 2.5, creativecommons.org/licenses/by-nc/2.5. Unchanged."], size=7, color=MUTED)
    # the words
    text(s, 3.3, 1.3, 9.45, 3.4, [
        "Most of what a product ships is code its own team never wrote. Dependencies move, advisories land, and the tests that would prove a change is safe never get written, so the gate runs on trust.",
        "The AI Test Harness writes those tests, runs them sealed on both sides of every change, and puts the evidence on the pull request: a signed record anyone can verify, an answer for every known vulnerability, and a plain account of what is proven, what is handed to you, and what is still a gap.",
    ], size=16, space_after=12)
    # where it sits in the delivery flow
    rule = s.shapes.add_connector(1, Inches(L), Inches(5.0), Inches(L + W), Inches(5.0)); rule.line.color.rgb = MID; rule.line.width = Pt(0.75)
    text(s, L, 5.08, W, 0.3, [[("Where it sits in your delivery flow", True)]], size=12, color=RED)
    steps = [("Code", None, 1), "→", ("Pull request", None, 1), "→", ("Your CI", None, 1), "+",
             ("AI Test Harness", "tests both versions, signs the evidence", 1.6), "→",
             ("Gate", "decides on evidence; either side can say no", 1.5), "→", ("Merge", None, 1), "→",
             ("Build and release", "signed records and VEX travel with it", 1.3)]
    conn_w = 0.25
    unit = (W - conn_w * sum(1 for x in steps if isinstance(x, str))) / sum(x[2] for x in steps if not isinstance(x, str))
    x, y, h = L, 5.45, 0.85
    for st in steps:
        if isinstance(st, str):
            text(s, x, y, conn_w, h, [st], size=12, color=MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=0)
            x += conn_w
            continue
        name, note, weight = st
        w = unit * weight
        hot, gate = name == "AI Test Harness", name == "Gate"
        lines = [(name, 11, True, WHITE if hot else INK, False)]
        if note:
            lines.append((note, 9, False, PINK if hot else MUTED, False))
        box(s, x, y, w, h, lines, fill=RED if hot else GREY, line=RED if hot else (INK if gate else LINE), line_w=1.5 if gate else 0.75)
        x += w
    text(s, L, 6.4, W, 0.3, ["An advisory or a dependency bump enters the same flow as a change, so a fix is proven the same way a feature is."], size=9, color=MUTED)


def slide_two(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "AI Test Harness")
    text(s, L, 2.3, 3.4, 0.9, [[("From a change to evidence at the gate", True)]], size=18, color=RED)
    text(s, L, 3.25, 3.4, 3.0, ["The harness reads a change and everything it touches, writes the tests nobody will write by hand, and runs them sealed on both versions. "
                                "Every behaviour the change reaches ends in one of three states. The evidence lands on the pull request, where a person and the gate decide."], size=14)
    x, w, y = 4.5, 8.23, 1.25
    band(s, x, y, w, "What it reads", red=True); y += 0.3
    y += grid(s, x, y, w, [["1. Your change", "2. Your dependencies", "3. Known vulnerabilities", "4. Your existing tests"],
                           ["code before and after", "every package, both versions", "public advisories", "what is already covered"],
                           ["what your code actually uses", "how deep the change goes", "what is already known to be wrong", "what your pipeline can run"]], row_h=0.4, size=9)
    down(s, x, y, w); y += 0.2
    band(s, x, y, w, "What it does", red=False); y += 0.3
    y += grid(s, x, y, w, [["Writes tests from facts, never guesses", "Breaks the code on purpose to see if they notice", "Runs them sealed, both versions, several times"]], row_h=0.5, size=9)
    down(s, x, y, w); y += 0.2
    band(s, x, y, w, "Every behaviour the change reaches ends in one of three states", red=False); y += 0.3
    states = [("decode()", "the test catches the breakage", "Validated", GREEN), ("upload()", "needs your network, not the sealed box", "Handed to you", INK),
              ("native()", "no test survived scoring", "Named gap, with its reason", MUTED)]
    rows = [[[(f"[ {fn} ]", 9, False, INK, False)], [("⟹  " + what.upper(), 9, False, INK, False)], [("● " + st.upper(), 9, True, col, False)]] for fn, what, st, col in states]
    y += grid(s, x, y, w, rows, col_w=[1.4, 3.9, w - 5.3], row_h=0.3, header=False, last_right=True)
    down(s, x, y, w); y += 0.2
    band(s, x, y, w, "What you get, on the pull request", red=True); y += 0.3
    grid(s, x, y, w, [["The tests and their evidence", "An answer per known vulnerability", "A signed record anyone can verify"],
                      [[("a person decides to merge", 9, False, MUTED, True)], [("a draft VEX your security team confirms", 9, False, MUTED, True)], [("in-toto, SLSA, OpenVEX", 9, False, MUTED, True)]]], row_h=0.4, size=9)


def main(out: Path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide_one(prs); slide_two(prs)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(out)


if __name__ == "__main__":
    main(Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[1] / "harness" / "out" / "two-slides.pptx")
